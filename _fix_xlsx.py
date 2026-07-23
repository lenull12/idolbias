import openpyxl
path = r"C:\Users\Admin\Documents\IdolBias_Roster_Master.xlsx"
wb = openpyxl.load_workbook(path)
ws = wb["Characters"]
fixed = 0
for r in range(2, 17):
    cell = ws.cell(row=r, column=16)
    old = cell.value
    if isinstance(old, str) and "VLOOKUP" in old:
        new = old.replace("$A$3:$B$10", "$A$2:$B$9")
        cell.value = new
        fixed += 1
# aussi corriger la faute de frappe archetype L3 "tête" -> "tete" (canonique)
c3 = ws.cell(row=3, column=9)
if c3.value and "tête" in c3.value:
    c3.value = c3.value.replace("tête", "tete")
    print("L3 archetype normalise: tete")
wb.save(path)
print(f"VLOOKUP fixed sur {fixed} lignes. Range -> $A$2:$B$9")
