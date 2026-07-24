"""Fill Japanese players: archetypes, style, and stats in Excel master."""
import openpyxl, re, random
random.seed(42)

EXCEL = r'C:\Users\Admin\Documents\IdolBias_Roster_Master.xlsx'
wb = openpyxl.load_workbook(EXCEL)
ws = wb['Characters']
st = wb['Stats & OVR']

def slug(name):
    s = name.lower().strip()
    rep = {'í':'i','é':'e','è':'e','á':'a','ñ':'n','ú':'u','ó':'o','ç':'c','à':'a','ä':'a','ö':'o','ü':'u','ō':'o'}
    for k,v in rep.items(): s = s.replace(k,v)
    s = re.sub(r'[^a-z0-9]+','-',s).strip('-')
    return s

# Archetypes + style per player
props = {
    'karen-himekami':    ('vitesse, finition, tete', 'percussion'),
    'shiori-saonji':     ('creatif, vision, trequarti', 'vista'),
    'reika-shinomiya':   ('vitesse, creatif', 'percussion'),
    'miyabi-kirishima':  ('vitesse, creatif, aerien', 'vista'),
    'hina-tsukiyomi':    ('vision, regista', 'vista'),
    'hana-kamishiro':    ('rugueux, pressing, mur', 'pressing'),
    'momo-hasegawa':     ('fullback, vision', 'vista'),
    'rin-morishita':     ('fullback, mur', 'pressing'),
    'yuriko-otake':      ('patron, mur, vision', 'elevation'),
    'aya-mishima':       ('patron, tete, mur', 'elevation'),
    'hinata-shigaki':    ('gardien', 'sangFroid'),
}

KEY_MAP = {
    'Vitesse':'vitesse','Accélération':'acceleration','Endurance':'endurance',
    'Puissance':'puissance','Agilité':'agilite','Détente':'detente',
    'Anticipation':'anticipation','SangFroid':'sangFroid','Leadership':'leadership',
    'Positionnement':'positionnement','Agressivité':'agressivite','Décision':'decision',
    'Passe':'passe','Tir':'tir','Dribble':'dribble','Centre':'centre',
    'Tacle':'tacle','Technique':'controle',
    'CF':'cf','Corners':'corners','Penalty':'penalty','LongThrows':'longThrows',
    'Réflexes':'reflexes','Handling':'handling','AerialReach':'aerialReach',
    'CommandArea':'commandArea','Kicking':'kicking','RushingOut':'rushingOut',
}
OUT_NAMES = ['Vitesse','Accélération','Endurance','Puissance','Agilité','Détente',
    'Anticipation','SangFroid','Leadership','Positionnement','Agressivité','Décision',
    'Passe','Tir','Dribble','Centre','Tacle','Technique',
    'CF','Corners','Penalty','LongThrows']
GK_NAMES = ['Vitesse','Accélération','Endurance','Puissance','Agilité','Détente',
    'Anticipation','SangFroid','Leadership','Positionnement','Agressivité','Décision',
    'Réflexes','Handling','AerialReach','CommandArea','Kicking','RushingOut']


def gen_out_stats(base, poste, role, taille, poids):
    """Generate 22 outfield stats based on position, role, height/weight."""
    lvl = base

    # Define emphases by (position, role)
    profiles = {
        ('ST','advanced_forward'): 
            (['tir','dribble','vitesse','acceleration','detente'],
             ['agilite','decision','sangFroid','controle','penalty','passe','centre'],
             ['tacle','longThrows','anticipation']),
        ('ST','poacher'): 
            (['tir','dribble','vitesse','acceleration','detente'],
             ['agilite','decision','sangFroid','controle','penalty'],
             ['passe','tacle','centre','longThrows','anticipation']),
        ('ST','second_striker'): 
            (['tir','dribble','passe','decision','controle'],
             ['vitesse','acceleration','agilite','sangFroid','cf'],
             ['tacle','puissance','centre','longThrows']),
        ('CAM','playmaker'): 
            (['passe','dribble','controle','decision','technique'],
             ['vitesse','sangFroid','agilite','tir'],
             ['tacle','agressivite','puissance','detente','longThrows']),
        ('CAM','trequartista'): 
            (['dribble','technique','passe','decision','controle'],
             ['agilite','tir','sangFroid','vitesse'],
             ['tacle','endurance','agressivite','longThrows']),
        ('LW','winger'): 
            (['centre','dribble','vitesse','acceleration','passe'],
             ['agilite','controle','endurance','decision'],
             ['tacle','puissance','tir','anticipation','positionnement']),
        ('RW','winger'): 
            (['centre','dribble','vitesse','acceleration','passe'],
             ['agilite','controle','endurance','decision'],
             ['tacle','puissance','tir','anticipation','positionnement']),
        ('CDM','regista'): 
            (['passe','decision','controle','anticipation'],
             ['tacle','positionnement','endurance','sangFroid','longThrows'],
             ['tir','dribble','centre','vitesse','detente']),
        ('CDM','ball_winning'): 
            (['tacle','agressivite','endurance','puissance','anticipation'],
             ['positionnement','decision','vitesse','agilite','passe'],
             ['tir','centre','dribble','technique']),
        ('LB','inverted_wingback'): 
            (['passe','centre','endurance','decision','controle'],
             ['vitesse','anticipation','tacle','agilite'],
             ['tir','detente','dribble','puissance']),
        ('RB','inverted_wingback'): 
            (['passe','centre','endurance','decision','controle'],
             ['vitesse','anticipation','tacle','agilite'],
             ['tir','detente','dribble','puissance']),
        ('LB','fullback'): 
            (['tacle','vitesse','endurance','positionnement','centre'],
             ['anticipation','passe','agilite','decision'],
             ['tir','dribble','detente','puissance','penalty']),
        ('RB','fullback'): 
            (['tacle','vitesse','endurance','positionnement','centre'],
             ['anticipation','passe','agilite','decision'],
             ['tir','dribble','detente','puissance','penalty']),
        ('CB','libero'): 
            (['passe','anticipation','positionnement','decision','controle'],
             ['tacle','agilite','sangFroid','vitesse'],
             ['tir','dribble','centre','detente','penalty']),
        ('CB','ball_playing_defender'): 
            (['tacle','positionnement','anticipation','puissance','passe'],
             ['decision','agressivite','detente','endurance'],
             ['tir','dribble','centre','vitesse','penalty','cf']),
        ('CB','stopper'): 
            (['tacle','puissance','positionnement','agressivite','anticipation'],
             ['detente','decision','endurance','vitesse'],
             ['passe','tir','dribble','centre','technique']),
    }

    key_stats, imp_stats, weak_stats = profiles.get((poste, role), 
        (['passe','decision','endurance'], ['controle','vitesse'], ['tir','tacle']))

    # Very weak stats = carrément inutiles pour ce poste (ex: TAC pour ST, TIR pour CB)
    if poste in ('ST', 'LW', 'RW', 'LM', 'RM'):
        very_weak = ['tacle']
        weak_stats = [s for s in weak_stats if s != 'tacle']
    elif poste in ('CB',):
        very_weak = ['tir']
        weak_stats = [s for s in weak_stats if s != 'tir']
    elif poste in ('LB', 'RB'):
        very_weak = ['tir']
        weak_stats = [s for s in weak_stats if s != 'tir']
    elif poste in ('CAM',):
        very_weak = ['tacle', 'agressivite']
        weak_stats = [s for s in weak_stats if s not in ('tacle', 'agressivite')]
    elif poste in ('CDM',):
        very_weak = ['tir', 'dribble']
        weak_stats = [s for s in weak_stats if s not in ('tir', 'dribble')]
    else:
        very_weak = []

    all_out = ['vitesse','acceleration','endurance','puissance','agilite','detente',
               'anticipation','sangFroid','leadership','positionnement','agressivite','decision',
               'passe','tir','dribble','centre','tacle','controle',
               'cf','corners','penalty','longThrows']

    stats = {}
    for s in all_out:
        if s in key_stats:
            val = min(99, lvl + random.randint(-2, 4))
            val = max(val, lvl - 3)
        elif s in imp_stats:
            val = min(99, lvl + random.randint(-5, 0))
        elif s in very_weak:
            val = min(99, max(15, lvl - random.randint(25, 40)))
        elif s in weak_stats:
            val = min(99, max(20, lvl - random.randint(15, 28)))
        else:
            val = min(99, max(25, lvl - random.randint(8, 16)))
        stats[s] = val

    # Height/weight adjustments
    if taille and taille >= 178:
        stats['detente'] = min(99, stats.get('detente', 50) + 4)
        stats['vitesse'] = max(40, stats.get('vitesse', 50) - 2)
    if taille and taille <= 162:
        stats['agilite'] = min(99, stats.get('agilite', 50) + 3)
        stats['vitesse'] = min(99, stats.get('vitesse', 50) + 2)
    if poids and poids >= 70:
        stats['puissance'] = min(99, stats.get('puissance', 50) + 3)
    if poids and poids <= 55:
        stats['agilite'] = min(99, stats.get('agilite', 50) + 2)

    return stats


def gen_gk_stats(base, taille, poids):
    """Generate 18 GK stats."""
    key_stats = ['reflexes', 'handling', 'anticipation', 'positionnement', 'rushingOut']
    imp_stats = ['aerialReach', 'commandArea', 'kicking', 'agilite', 'decision', 'sangFroid']
    weak_stats = ['vitesse', 'endurance', 'agressivite', 'leadership', 'puissance']
    rest = ['detente', 'acceleration']

    all_gk = ['vitesse','acceleration','endurance','puissance','agilite','detente',
              'anticipation','sangFroid','leadership','positionnement','agressivite','decision',
              'reflexes','handling','aerialReach','commandArea','kicking','rushingOut']

    stats = {}
    for s in all_gk:
        if s in key_stats:
            val = min(99, base + random.randint(-2, 4))
            val = max(val, base - 3)
        elif s in imp_stats:
            val = min(99, base + random.randint(-5, 0))
        elif s in weak_stats:
            val = min(99, max(30, base - random.randint(8, 18)))
        else:
            val = min(99, base - random.randint(4, 10))
            val = max(val, 40)
        stats[s] = val

    if taille and taille >= 175:
        stats['aerialReach'] = min(99, stats.get('aerialReach', 50) + 5)
    if taille and taille <= 165:
        stats['agilite'] = min(99, stats.get('agilite', 50) + 3)
        stats['reflexes'] = min(99, stats.get('reflexes', 50) + 2)

    return stats


# ─── Main loop ───
print('=== Stats générées pour le Japon ===')
for r in range(14, 25):
    name = str(ws.cell(row=r, column=1).value or '').strip()
    nation = str(ws.cell(row=r, column=3).value or '').strip()
    if nation != 'japon' or not name:
        continue

    poste = str(ws.cell(row=r, column=4).value or '').strip().upper()
    role = str(ws.cell(row=r, column=5).value or '').strip()
    base_val = ws.cell(row=r, column=10).value
    if isinstance(base_val, str):
        base_val = int(float(base_val))
    taille = ws.cell(row=r, column=12).value
    poids = ws.cell(row=r, column=13).value
    sid = slug(name)
    is_gk = (poste == 'GK')

    # Generate stats
    if is_gk:
        stats = gen_gk_stats(base_val, taille, poids)
    else:
        stats = gen_out_stats(base_val, poste, role, taille, poids)

    # Show summary
    key_vals = [f'{k}={stats[k]}' for k in list(stats.keys())[:6]]
    print(f'  {name:25s} ({poste:3s} base={base_val:2d}) {sid:20s}', end='')
    print(f'  {", ".join(key_vals)}...')

    # Write archetypes + style to Characters sheet
    arches, style = props.get(sid, ('', ''))
    ws.cell(row=r, column=8).value = arches
    ws.cell(row=r, column=9).value = style

    # Write to Stats & OVR sheet (same row)
    st.cell(row=r, column=1).value = sid
    st.cell(row=r, column=2).value = name
    st.cell(row=r, column=3).value = nation
    st.cell(row=r, column=4).value = poste
    st.cell(row=r, column=5).value = base_val

    col_start = 31 if is_gk else 6
    names = GK_NAMES if is_gk else OUT_NAMES
    for i, n in enumerate(names):
        sk = KEY_MAP[n]
        val = stats.get(sk, 0)
        st.cell(row=r, column=col_start + i).value = val

wb.save(EXCEL)
print()
print(f'✅ Fini ! Excel mis à jour avec 11 joueuses japonaises.')
