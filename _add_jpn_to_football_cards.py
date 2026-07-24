"""Add missing Japanese characters to footballCards.ts."""
import openpyxl, re

EXCEL = r'C:\Users\Admin\Documents\IdolBias_Roster_Master.xlsx'
wb = openpyxl.load_workbook(EXCEL, data_only=True)
ws = wb['Characters']
TS_PATH = r'//wsl.localhost/Ubuntu/home/raphi/idolbias/src/data/footballCards.ts'

def slug(name):
    s = name.lower().strip()
    rep = {'í':'i','é':'e','è':'e','á':'a','ñ':'n','ú':'u','ó':'o','ç':'c','à':'a','ä':'a','ö':'o','ü':'u','ō':'o'}
    for k,v in rep.items(): s = s.replace(k,v)
    s = re.sub(r'[^a-z0-9]+','-',s).strip('-')
    return s

# Read Japanese rows
jpn_entries = []
for r in range(14, 25):
    name = str(ws.cell(row=r, column=1).value or '').strip()
    nation = str(ws.cell(row=r, column=3).value or '').strip()
    if nation != 'japon' or not name: continue
    
    sid = slug(name)
    style = str(ws.cell(row=r, column=9).value or '').strip()
    poste = str(ws.cell(row=r, column=4).value or '').strip().upper()
    is_captain = str(ws.cell(row=r, column=14).value or '').strip().lower() == 'oui'
    
    # Map position12 to group
    pos_map = {'GK':'GB','RB':'DEF','LB':'DEF','CB':'DEF','CDM':'MIL','CM':'MIL','CAM':'MIL',
               'LM':'ATT','RM':'ATT','LW':'ATT','RW':'ATT','ST':'ATT'}
    group = pos_map.get(poste, 'ATT')
    
    entry = f"""  {{
    id: "{sid}",
    name: "{name}",
    nation: "japon",
    defaultStyle: "{style}",
    defaultPosition: "{group}",
    isCaptain: {"true" if is_captain else "false"},
    photoVariants: {{
      standard: "/cards/football/japon/{sid}/standard.png",
    }},
  }},"""
    jpn_entries.append(entry)
    print(f'  ✅ {sid}')

# Find the insertion point in the TS file: end of CHARACTERS array
with open(TS_PATH, 'r', encoding='utf-8') as f:
    content = f.read()

# Find the last character entry before the closing ]
# The pattern is: esperanza-galvan is the last Argentine, then ]};
insert_marker = '];\n\nexport const CARD_PRINTS'
insert_pos = content.find(insert_marker)

if insert_pos == -1:
    print("Could not find insertion point!")
    exit(1)

# Build the block to insert
block = '\n' + '\n'.join(jpn_entries) + '\n'

new_content = content[:insert_pos] + block + content[insert_pos:]

with open(TS_PATH, 'w', encoding='utf-8') as f:
    f.write(new_content)

print(f'\n✅ Added {len(jpn_entries)} Japanese characters to CHARACTERS')
