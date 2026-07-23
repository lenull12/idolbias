import openpyxl
path = r"C:\Users\Admin\Documents\IdolBias_Roster_Master.xlsx"
wb = openpyxl.load_workbook(path, data_only=False)
ws = wb["Characters"]
print("Feuille Characters — lignes max:", ws.max_row)
hdr = [c.value for c in ws[1]]
print("HEADERS:", hdr[:11])
print("---")
rows = []
for r in range(2, ws.max_row+1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 17)]
    if any(v not in (None, "") for v in vals):
        rows.append((r, vals))
print(len(rows), "lignes remplies (hors entete)")
for r, v in rows:
    print(f"L{r}: id={v[0]} | {v[1]} | nat={v[3]} | poste={v[4]} | role={v[5]} | sec={v[6]},{v[7]} | arch={v[8]} | style={v[9]} | ceil={v[10]} | cap={v[11]} | photo={v[12]} | num={v[14]} | prefix={v[15]}")
# Stats & OVR preview
st = wb["Stats & OVR"]
print("\n--- Stats & OVR (ceiling auto) ---")
for r in range(2, st.max_row+1):
    cid = st.cell(row=r, column=1).value
    if cid not in (None, ""):
        ceil = st.cell(row=r, column=5).value
        ovr = st.cell(row=r, column=28).value
        tier = st.cell(row=r, column=29).value
        print(f"L{r}: {cid} | ceiling={ceil} | ovr={ovr} | tier={tier}")
